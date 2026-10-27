import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_para_rtl(p, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p.alignment = align
    pPr = p._element.get_or_add_pPr()
    bidi = pPr.find(qn('w:bidi'))
    if bidi is None:
        bidi = OxmlElement('w:bidi')
        pPr.append(bidi)
    bidi.set(qn('w:val'), '1')

def add_run(p, text, font_name="B Nazanin", size_pt=10, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = font_name
    r.font.size = Pt(size_pt)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rtl = rPr.find(qn('w:rtl'))
    if rtl is None:
        rtl = OxmlElement('w:rtl')
        rPr.append(rtl)
    rtl.set(qn('w:val'), '1')
    return r

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_styling(table, border_color="CBD5E1", border_sz="4"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'  <w:left w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'  <w:bottom w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'  <w:right w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'  <w:insideH w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'  <w:insideV w:val="single" w:sz="{border_sz}" w:space="0" w:color="{border_color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_header_table_borders(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tblPr = table._tbl.tblPr
    tblPr.append(parse_xml(f'<w:bidiVisual {nsdecls("w")}/>'))
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="0F172A"/>'
        f'  <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideH w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'  <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def build_template():
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        sectPr = section._sectPr
        bidi = sectPr.find(qn('w:bidi'))
        if bidi is None:
            bidi = OxmlElement('w:bidi')
            sectPr.append(bidi)
        bidi.set(qn('w:val'), '1')

    header_table = doc.add_table(rows=1, cols=3)
    header_table.autofit = False
    set_header_table_borders(header_table)
    
    col_widths = [Inches(1.8), Inches(3.6), Inches(1.8)]
    for row in header_table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = width

    cell_logo = header_table.cell(0, 0)
    cell_logo.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_logo = cell_logo.paragraphs[0]
    p_logo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    logo_path = os.path.join(os.path.dirname(__file__), "..", "templates", "Logo.png")
    if os.path.exists(logo_path):
        p_logo.add_run().add_picture(logo_path, width=Inches(1.2))

    cell_center = header_table.cell(0, 1)
    cell_center.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_c1 = cell_center.paragraphs[0]
    set_para_rtl(p_c1, WD_ALIGN_PARAGRAPH.CENTER)
    p_c1.paragraph_format.space_after = Pt(2)
    add_run(p_c1, "« بسمه تعالی »", font_name="B Nazanin", size_pt=11, bold=True, color="334155")
    
    p_c2 = cell_center.add_paragraph()
    set_para_rtl(p_c2, WD_ALIGN_PARAGRAPH.CENTER)
    p_c2.paragraph_format.space_after = Pt(2)
    add_run(p_c2, "برگ پرونده گزارش محرمانه", font_name="B Titr", size_pt=14, bold=True, color="0F172A")
    
    p_c3 = cell_center.add_paragraph()
    set_para_rtl(p_c3, WD_ALIGN_PARAGRAPH.CENTER)
    p_c3.paragraph_format.space_after = Pt(0)
    add_run(p_c3, "سامانه ستاد خبری و گزارش‌های مردمی", font_name="B Nazanin", size_pt=10, bold=False, color="475569")

    cell_meta = header_table.cell(0, 2)
    cell_meta.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p_m1 = cell_meta.paragraphs[0]
    set_para_rtl(p_m1, WD_ALIGN_PARAGRAPH.LEFT)
    p_m1.paragraph_format.space_after = Pt(2)
    add_run(p_m1, "شماره پرونده: ", font_name="B Nazanin", size_pt=9, bold=True)
    add_run(p_m1, "{{TRACKING_CODE}}", font_name="Calibri", size_pt=10, bold=True, color="0284C7")
    
    p_m2 = cell_meta.add_paragraph()
    set_para_rtl(p_m2, WD_ALIGN_PARAGRAPH.LEFT)
    p_m2.paragraph_format.space_after = Pt(2)
    add_run(p_m2, "تاریخ ثبت: ", font_name="B Nazanin", size_pt=9, bold=True)
    add_run(p_m2, "{{DATE}}", font_name="B Nazanin", size_pt=9)

    p_m3 = cell_meta.add_paragraph()
    set_para_rtl(p_m3, WD_ALIGN_PARAGRAPH.LEFT)
    p_m3.paragraph_format.space_after = Pt(2)
    add_run(p_m3, "زمان چاپ: ", font_name="B Nazanin", size_pt=9, bold=True)
    add_run(p_m3, "{{PRINT_DATE}}", font_name="B Nazanin", size_pt=9)

    p_m4 = cell_meta.add_paragraph()
    set_para_rtl(p_m4, WD_ALIGN_PARAGRAPH.LEFT)
    p_m4.paragraph_format.space_after = Pt(0)
    add_run(p_m4, "طبقه‌بندی: ", font_name="B Nazanin", size_pt=9, bold=True)
    add_run(p_m4, "محرمانه", font_name="B Nazanin", size_pt=9, bold=True, color="DC2626")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    meta_table = doc.add_table(rows=3, cols=4)
    meta_table.autofit = False
    set_table_styling(meta_table, border_color="CBD5E1", border_sz="4")
    
    meta_widths = [Inches(1.3), Inches(2.3), Inches(1.3), Inches(2.3)]
    for row in meta_table.rows:
        for idx, width in enumerate(meta_widths):
            row.cells[idx].width = width

    cells_config = [
        [(0, 0, "موضوع گزارش:", True, "F1F5F9"), (0, 1, "{{CATEGORY}}", False, "FFFFFF"), (0, 2, "اولویت رسیدگی:", True, "F1F5F9"), (0, 3, "{{PRIORITY}}", False, "FFFFFF")],
        [(1, 0, "نوع گزارش:", True, "F1F5F9"), (1, 1, "{{REPORTER_TYPE}}", False, "FFFFFF"), (1, 2, "مشخصات گزارش‌دهنده:", True, "F1F5F9"), (1, 3, "{{REPORTER_INFO}}", False, "FFFFFF")],
        [(2, 0, "مستندات پیوست:", True, "F1F5F9"), (2, 1, "{{EVIDENCE_STATUS}}", False, "FFFFFF"), (2, 2, "وضعیت بررسی پرونده:", True, "F1F5F9"), (2, 3, "{{READ_STATUS}}", False, "FFFFFF")]
    ]
    for row_conf in cells_config:
        for r_idx, c_idx, text, bold, bg in row_conf:
            cell = meta_table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
            p = cell.paragraphs[0]
            set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
            p.paragraph_format.space_after = Pt(0)
            add_run(p, text, font_name="B Nazanin", size_pt=9.5, bold=bold)

    def add_section_header(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        set_para_rtl(p, WD_ALIGN_PARAGRAPH.RIGHT)
        add_run(p, title_text, font_name="B Titr", size_pt=10.5, bold=True, color="0F172A")

    add_section_header("۱. متن و شرح گزارش واصله:")
    tbl_desc = doc.add_table(rows=1, cols=1)
    tbl_desc.autofit = False
    tbl_desc.columns[0].width = Inches(7.2)
    set_table_styling(tbl_desc, border_color="CBD5E1", border_sz="4")
    cell_desc = tbl_desc.cell(0, 0)
    set_cell_margins(cell_desc, top=100, bottom=100, left=120, right=120)
    p_desc = cell_desc.paragraphs[0]
    set_para_rtl(p_desc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_desc.paragraph_format.space_after = Pt(0)
    p_desc.paragraph_format.line_spacing = 1.25
    add_run(p_desc, "{{DESCRIPTION}}", font_name="B Nazanin", size_pt=10)

    add_section_header("۲. پاسخ رسمی و اقدامات صورت‌گرفته توسط حراست:")
    tbl_resp = doc.add_table(rows=2, cols=2)
    tbl_resp.autofit = False
    set_table_styling(tbl_resp, border_color="CBD5E1", border_sz="4")
    
    cell_resp_content = tbl_resp.cell(0, 0)
    cell_resp_content.merge(tbl_resp.cell(0, 1))
    set_cell_margins(cell_resp_content, top=100, bottom=100, left=120, right=120)
    p_rc = cell_resp_content.paragraphs[0]
    set_para_rtl(p_rc, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_rc.paragraph_format.space_after = Pt(0)
    p_rc.paragraph_format.line_spacing = 1.25
    add_run(p_rc, "{{ADMIN_RESPONSE}}", font_name="B Nazanin", size_pt=10)

    cell_by = tbl_resp.cell(1, 0)
    cell_by.width = Inches(3.6)
    set_cell_shading(cell_by, "F8FAFC")
    set_cell_margins(cell_by, top=60, bottom=60, left=100, right=100)
    p_by = cell_by.paragraphs[0]
    set_para_rtl(p_by, WD_ALIGN_PARAGRAPH.RIGHT)
    p_by.paragraph_format.space_after = Pt(0)
    add_run(p_by, "ثبت پاسخ توسط: ", font_name="B Nazanin", size_pt=9, bold=True)
    add_run(p_by, "{{RESPONDED_BY}}", font_name="B Nazanin", size_pt=9)

    cell_at = tbl_resp.cell(1, 1)
    cell_at.width = Inches(3.6)
    set_cell_shading(cell_at, "F8FAFC")
    set_cell_margins(cell_at, top=60, bottom=60, left=100, right=100)
    p_at = cell_at.paragraphs[0]
    set_para_rtl(p_at, WD_ALIGN_PARAGRAPH.LEFT)
    p_at.paragraph_format.space_after = Pt(0)
    add_run(p_at, "تاریخ ثبت پاسخ: ", font_name="B Nazanin", size_pt=9, bold=True)
    add_run(p_at, "{{RESPONDED_AT}}", font_name="B Nazanin", size_pt=9)

    add_section_header("۳. ملاحظات و یادداشت‌های محرمانه داخلی پرونده:")
    tbl_notes = doc.add_table(rows=1, cols=1)
    tbl_notes.autofit = False
    tbl_notes.columns[0].width = Inches(7.2)
    set_table_styling(tbl_notes, border_color="CBD5E1", border_sz="4")
    cell_notes = tbl_notes.cell(0, 0)
    set_cell_margins(cell_notes, top=100, bottom=100, left=120, right=120)
    p_notes = cell_notes.paragraphs[0]
    set_para_rtl(p_notes, WD_ALIGN_PARAGRAPH.JUSTIFY)
    p_notes.paragraph_format.space_after = Pt(0)
    p_notes.paragraph_format.line_spacing = 1.25
    add_run(p_notes, "{{INTERNAL_NOTES}}", font_name="B Nazanin", size_pt=10)

    add_section_header("۴. صورت‌جلسه و تأییدیه نهایی بررسی پرونده:")
    tbl_sig = doc.add_table(rows=1, cols=2)
    tbl_sig.autofit = False
    set_table_styling(tbl_sig, border_color="CBD5E1", border_sz="4")
    
    cell_s1 = tbl_sig.cell(0, 0)
    cell_s1.width = Inches(3.6)
    set_cell_margins(cell_s1, top=100, bottom=100, left=120, right=120)
    p_s1 = cell_s1.paragraphs[0]
    set_para_rtl(p_s1, WD_ALIGN_PARAGRAPH.CENTER)
    p_s1.paragraph_format.space_after = Pt(4)
    add_run(p_s1, "کارشناس بررسی‌کننده پرونده", font_name="B Nazanin", size_pt=10, bold=True)
    
    p_s1_sub = cell_s1.add_paragraph()
    set_para_rtl(p_s1_sub, WD_ALIGN_PARAGRAPH.RIGHT)
    p_s1_sub.paragraph_format.space_after = Pt(28)
    add_run(p_s1_sub, "نام و سمت سازمانی: .......................................", font_name="B Nazanin", size_pt=9)
    
    p_s1_ft = cell_s1.add_paragraph()
    set_para_rtl(p_s1_ft, WD_ALIGN_PARAGRAPH.CENTER)
    p_s1_ft.paragraph_format.space_after = Pt(0)
    add_run(p_s1_ft, "محل امضاء و تاریخ", font_name="B Nazanin", size_pt=9, bold=True, color="64748B")

    cell_s2 = tbl_sig.cell(0, 1)
    cell_s2.width = Inches(3.6)
    set_cell_margins(cell_s2, top=100, bottom=100, left=120, right=120)
    p_s2 = cell_s2.paragraphs[0]
    set_para_rtl(p_s2, WD_ALIGN_PARAGRAPH.CENTER)
    p_s2.paragraph_format.space_after = Pt(4)
    add_run(p_s2, "رئیس / مدیر اداره حراست", font_name="B Nazanin", size_pt=10, bold=True)
    
    p_s2_sub = cell_s2.add_paragraph()
    set_para_rtl(p_s2_sub, WD_ALIGN_PARAGRAPH.RIGHT)
    p_s2_sub.paragraph_format.space_after = Pt(28)
    add_run(p_s2_sub, "نام و سمت سازمانی: .......................................", font_name="B Nazanin", size_pt=9)
    
    p_s2_ft = cell_s2.add_paragraph()
    set_para_rtl(p_s2_ft, WD_ALIGN_PARAGRAPH.CENTER)
    p_s2_ft.paragraph_format.space_after = Pt(0)
    add_run(p_s2_ft, "محل امضاء و تاریخ", font_name="B Nazanin", size_pt=9, bold=True, color="64748B")

    p_warn = doc.add_paragraph()
    p_warn.paragraph_format.space_before = Pt(14)
    p_warn.paragraph_format.space_after = Pt(0)
    set_para_rtl(p_warn, WD_ALIGN_PARAGRAPH.CENTER)
    add_run(p_warn, "⚠️ هشدار امنیتی: این سند حاوی اطلاعات طبقه‌بندی‌شده و محرمانه اداری است. هرگونه نسخه‌برداری غیرمجاز یا افشای مفاد آن پیگرد قانونی به همراه دارد.", font_name="B Nazanin", size_pt=8.5, bold=True, color="B91C1C")

    out_path = os.path.join(os.path.dirname(__file__), "..", "templates", "report_template.docx")
    doc.save(out_path)

if __name__ == "__main__":
    build_template()
