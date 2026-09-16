import sys 
import os 
import json 
import io 
import re 
import docx 
from docx .oxml import OxmlElement 
from docx .oxml .ns import qn 
from docx .enum .text import WD_ALIGN_PARAGRAPH 

def set_persian_formatting (run ,font_name ="B Nazanin"):
    rPr =run ._element .get_or_add_rPr ()


    rFonts =rPr .find (qn ('w:rFonts'))
    if rFonts is None :
        rFonts =OxmlElement ('w:rFonts')
        rPr .append (rFonts )
    rFonts .set (qn ('w:ascii'),font_name )
    rFonts .set (qn ('w:hAnsi'),font_name )
    rFonts .set (qn ('w:cs'),font_name )


    rtl =rPr .find (qn ('w:rtl'))
    if rtl is None :
        rtl =OxmlElement ('w:rtl')
        rPr .append (rtl )
    rtl .set (qn ('w:val'),'1')

def set_english_formatting (run ,font_name ="Calibri"):
    rPr =run ._element .get_or_add_rPr ()

    rFonts =rPr .find (qn ('w:rFonts'))
    if rFonts is None :
        rFonts =OxmlElement ('w:rFonts')
        rPr .append (rFonts )
    rFonts .set (qn ('w:ascii'),font_name )
    rFonts .set (qn ('w:hAnsi'),font_name )
    rFonts .set (qn ('w:cs'),font_name )


    rtl =rPr .find (qn ('w:rtl'))
    if rtl is None :
        rtl =OxmlElement ('w:rtl')
        rPr .append (rtl )
    rtl .set (qn ('w:val'),'0')

def format_persian_text_with_dates (text ):
    if not text :
        return text 

    text =re .sub (r'(\d{4}/\d{2}/\d{2})\s*-\s*(\d{2}:\d{2})',lambda m :'\u200E'+m .group (1 )+'\u200E - \u200E'+m .group (2 )+'\u200E',str (text ))

    text =re .sub (r'(?<!\u200E)(\d{4}/\d{2}/\d{2})(?!\u200E)',lambda m :'\u200E'+m .group (1 )+'\u200E',text )
    return text 

def replace_placeholders_in_paragraph (p ,mapping ):
    full_text =p .text 
    if not any (tag in full_text for tag in mapping ):
        return 


    for run in p .runs :
        for tag ,val in mapping .items ():
            if tag in run .text :
                if tag =="{{TRACKING_CODE}}":
                    run .text =run .text .replace (tag ,str (val ))
                    set_english_formatting (run ,font_name ="Calibri")
                else :
                    formatted_val =format_persian_text_with_dates (val )
                    run .text =run .text .replace (tag ,formatted_val )
                    set_persian_formatting (run ,font_name ="B Nazanin")


    if any (tag in p .text for tag in mapping ):
        merged_text =p .text 
        for tag ,val in mapping .items ():
            if tag in merged_text :
                if tag =="{{TRACKING_CODE}}":
                    merged_text =merged_text .replace (tag ,str (val ))
                else :
                    merged_text =merged_text .replace (tag ,format_persian_text_with_dates (val ))

        if p .runs :
            p .runs [0 ].text =merged_text 
            set_persian_formatting (p .runs [0 ],font_name ="B Nazanin")
            for r in p .runs [1 :]:
                r .text =""


    pPr =p ._element .get_or_add_pPr ()
    bidi =pPr .find (qn ('w:bidi'))
    if bidi is None :
        bidi =OxmlElement ('w:bidi')
        pPr .append (bidi )
    bidi .set (qn ('w:val'),'1')

def replace_in_table (table ,mapping ):
    for row in table .rows :
        for cell in row .cells :
            for p in cell .paragraphs :
                replace_placeholders_in_paragraph (p ,mapping )
            for sub_table in cell .tables :
                replace_in_table (sub_table ,mapping )

def populate_docx (template_path ,data ):
    if not os .path .exists (template_path ):
        raise FileNotFoundError (f"Template not found at: {template_path }")

    doc =docx .Document (template_path )

    mapping ={
    "{{TRACKING_CODE}}":data .get ("tracking_code",""),
    "{{DATE}}":data .get ("date",""),
    "{{PRINT_DATE}}":data .get ("print_date",""),
    "{{CATEGORY}}":data .get ("category",""),
    "{{PRIORITY}}":data .get ("priority","عادی"),
    "{{REPORTER_TYPE}}":data .get ("reporter_type","ناشناس"),
    "{{REPORTER_INFO}}":data .get ("reporter_info","محفوظ"),
    "{{EVIDENCE_STATUS}}":data .get ("evidence_status","فاقد پیوست"),
    "{{READ_STATUS}}":data .get ("read_status","در انتظار بررسی"),
    "{{DESCRIPTION}}":data .get ("description",""),
    "{{ADMIN_RESPONSE}}":data .get ("admin_response","فاقد پاسخ کارشناس حراست"),
    "{{RESPONDED_BY}}":data .get ("responded_by","ثبت نشده"),
    "{{RESPONDED_AT}}":data .get ("responded_at","ثبت نشده"),
    "{{INTERNAL_NOTES}}":data .get ("internal_notes","فاقد یادداشت محرمانه داخلی.")
    }


    for p in doc .paragraphs :
        replace_placeholders_in_paragraph (p ,mapping )


    for table in doc .tables :
        replace_in_table (table ,mapping )


    for s in doc .sections :
        for p in s .header .paragraphs :
            replace_placeholders_in_paragraph (p ,mapping )
        for t in s .header .tables :
            replace_in_table (t ,mapping )
        for p in s .footer .paragraphs :
            replace_placeholders_in_paragraph (p ,mapping )
        for t in s .footer .tables :
            replace_in_table (t ,mapping )

    buf =io .BytesIO ()
    doc .save (buf )
    return buf .getvalue ()

def main ():
    if len (sys .argv )>1 and sys .argv [1 ]=="--help":
        print ("Usage: python generate_docx.py [template_path]")
        sys .exit (0 )

    template_path =sys .argv [1 ]if len (sys .argv )>1 else os .path .join (os .path .dirname (__file__ ),"..","templates","report_template.docx")


    raw_input =sys .stdin .buffer .read ()
    if not raw_input .strip ():
        sys .stderr .write ("Error: No JSON input provided on stdin\n")
        sys .exit (1 )

    try :
        json_text =raw_input .decode ('utf-8')
        data =json .loads (json_text )
        docx_bytes =populate_docx (template_path ,data )

        sys .stdout .buffer .write (docx_bytes )
        sys .stdout .buffer .flush ()
    except Exception as e :
        sys .stderr .write (f"Error generating docx: {str (e )}\n")
        sys .exit (1 )

if __name__ =="__main__":
    main ()
